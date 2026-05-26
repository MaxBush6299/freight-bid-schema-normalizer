"""local_rehydrate_runner.py

Local entry point for the Rehydrate (reverse) pipeline.

Usage
-----
    python -m src.function_app.local_rehydrate_runner \\
        --template  "source_docs/Original Customer File 1.xlsx" \\
        --export    "source_docs/FTL LaneExport.xlsx" \\
        --output-root artifacts/local_rehydrate

    # LLM mapping with human-in-the-loop for low-confidence fields:
    python -m src.function_app.local_rehydrate_runner \\
        --template  "source_docs/Original Customer File 1.xlsx" \\
        --export    "source_docs/FTL LaneExport.xlsx" \\
        --planner-mode live \\
        --interactive \\
        --confidence-threshold 0.70

Outputs (written to <output-root>/<run_id>/)
--------------------------------------------
    submission.xlsx        — original template with pricing filled in
    write_report.json      — cells written / skipped / no-bid lanes
    template_profile.json  — TemplateProfile used for this run
    mapping_plan.json      — ReverseMappingPlan used for this run
    pending_review.json    — low-confidence mappings (empty list when mode=mock)
"""
from __future__ import annotations

import argparse
import json
import logging
import os
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models.contracts import (
    FieldMapping,
    HumanReviewRequest,
    ReverseValidationIssue,
    ReverseValidationReport,
    TemplateProfile,
)
from .services.reverse_planner import ReversePlanner
from .services.template_aware_writer import TemplateAwareWriter, resolve_instructions
from .services.template_diff_validator import TemplateDiffValidator
from .services.template_profiler import TemplateProfiler
from .services.xls_converter import ensure_xlsx

logger = logging.getLogger(__name__)

# P2-004: confidence at/above this score → mapping is eligible to write
# automatically. Lower-confidence mappings flow to pending_review but produce
# no writes unless explicitly approved by a human reviewer.
_DEFAULT_AUTO_WRITE_THRESHOLD = 0.85
_AUTO_WRITE_THRESHOLD_ENV = "REHYDRATE_MIN_WRITE_CONFIDENCE"


def _load_export_rows(export_path: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Load all rows from the priced export file as dicts keyed by header."""
    xlsx_path, _did_convert = ensure_xlsx(export_path)
    wb = load_workbook(str(xlsx_path), data_only=True, read_only=True)
    try:
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        records = [
            {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            for row in rows[1:]
            if any(v is not None for v in row)
        ]
        return records, headers
    finally:
        wb.close()


def _build_reverse_validation_report(
    no_bid_lanes: list[str],
    template_profile: TemplateProfile,
) -> ReverseValidationReport:
    """Build warning issues for template lanes missing from the priced export."""
    provenance_by_route = {}
    for sheet_profile in template_profile.bid_sheets:
        for entry in sheet_profile.lane_provenance:
            provenance_by_route.setdefault(entry.route_name, entry)
    warnings: list[ReverseValidationIssue] = []

    unique_no_bid_lanes = list(dict.fromkeys(no_bid_lanes))
    for route_name in unique_no_bid_lanes:
        provenance = provenance_by_route.get(route_name)
        warnings.append(ReverseValidationIssue(
            code="no_bid_lane",
            severity="warning",
            route_name=route_name,
            sheet_name=provenance.sheet_name if provenance else None,
            row_index=provenance.row_index if provenance else None,
            message=(
                "Template lane was not present in the priced export; "
                "submission cells were left blank for this route."
            ),
        ))

    return ReverseValidationReport(
        status="Passed",
        passed=True,
        issues=warnings,
        issue_counts={"error": 0, "warning": len(warnings)},
        no_bid_lanes=unique_no_bid_lanes,
    )


def _interactive_review(
    pending: list[HumanReviewRequest],
    mappings: list[FieldMapping],
) -> list[FieldMapping]:
    """Prompt the user to accept or override each low-confidence mapping.

    Updates the mapping list in-place and returns it.
    """
    if not pending:
        return mappings

    print("\n" + "=" * 60)
    print("HUMAN REVIEW REQUIRED")
    print(f"{len(pending)} mapping(s) have confidence below the review threshold.")
    print("For each, press Enter to accept the LLM suggestion, or type an override.\n")

    mapping_by_src = {fm.source_field: fm for fm in mappings}

    for req in pending:
        print(f"  Source field : {req.source_field!r}")
        print(f"  LLM target   : {req.target_column!r} (confidence {req.confidence_score:.0%})")
        if req.reasoning:
            print(f"  Reasoning    : {req.reasoning}")
        override = input("  Override target (or Enter to accept): ").strip()
        if override:
            req.override_value = override
            if req.source_field in mapping_by_src:
                mapping_by_src[req.source_field].target_column = override
                mapping_by_src[req.source_field].reasoning = f"human override: {override!r}"
            print(f"  ✓ Overridden → {override!r}\n")
        else:
            print("  ✓ Accepted\n")

    print("=" * 60 + "\n")
    return mappings


def _resolve_auto_write_threshold(explicit: float | None) -> float:
    """Resolve the auto-write threshold from (in priority order):
    explicit argument → ``REHYDRATE_MIN_WRITE_CONFIDENCE`` env var → default."""
    if explicit is not None:
        return float(explicit)
    raw = os.getenv(_AUTO_WRITE_THRESHOLD_ENV)
    if raw is None or raw.strip() == "":
        return _DEFAULT_AUTO_WRITE_THRESHOLD
    try:
        return float(raw)
    except ValueError:
        logger.warning(
            "Invalid %s=%r; falling back to default %.2f",
            _AUTO_WRITE_THRESHOLD_ENV,
            raw,
            _DEFAULT_AUTO_WRITE_THRESHOLD,
        )
        return _DEFAULT_AUTO_WRITE_THRESHOLD


def run_rehydrate(
    template_path: str,
    export_path: str,
    output_root: str,
    planner_mode: str = "mock",
    confidence_threshold: float = 0.70,
    interactive: bool = False,
    auto_write_threshold: float | None = None,
) -> dict[str, Any]:
    run_id = datetime.now(UTC).strftime("%Y%m%d%H%M%S")
    run_dir = Path(output_root) / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    auto_write_threshold_resolved = _resolve_auto_write_threshold(auto_write_threshold)

    # 1. Convert template .xls → .xlsx if needed
    template_xlsx, _did_convert = ensure_xlsx(template_path)

    # 2. Profile the template
    profiler = TemplateProfiler()
    template_profile = profiler.profile(str(template_xlsx))

    # 3. Load priced export rows
    export_rows, export_columns = _load_export_rows(export_path)

    # 4. Build mapping plan
    planner = ReversePlanner(mode=planner_mode, review_threshold=confidence_threshold)
    plan = planner.build_plan(template_profile, export_columns)

    # 5. Human-in-the-loop for low-confidence mappings
    if interactive and plan.pending_review:
        plan.mappings = _interactive_review(plan.pending_review, plan.mappings)

    # 6. Resolve CellWriteInstructions (gates by confidence + preserves filled cells +
    #    validates row descriptors). Mock plans bypass the confidence gate because
    #    mock mappings carry confidence 1.0 by construction.
    resolution = resolve_instructions(
        export_rows,
        template_profile,
        plan,
        template_path=str(template_xlsx),
        auto_write_threshold=auto_write_threshold_resolved,
        validate_descriptors=True,
    )
    instructions = resolution.instructions
    no_bid_lanes = resolution.no_bid_lanes

    # 7. Write submission
    submission_path = str(run_dir / "submission.xlsx")
    writer = TemplateAwareWriter()
    write_report = writer.write(
        template_path=str(template_xlsx),
        instructions=instructions,
        output_path=submission_path,
        template_profile=template_profile,
    )
    write_report.no_bid_lanes = no_bid_lanes
    write_report.cells_skipped_preserved_resolver = resolution.cells_skipped_preserved
    write_report.cells_skipped_low_confidence = resolution.cells_skipped_low_confidence
    write_report.rows_skipped_descriptor_mismatch = resolution.rows_skipped_descriptor_mismatch
    write_report.duplicate_export_routes = resolution.duplicate_export_routes
    # Merge resolver skip log into write_log so the artifact carries one timeline.
    write_report.write_log.extend(resolution.skip_log)
    write_report.validation_summary = _build_reverse_validation_report(no_bid_lanes, template_profile)
    write_report.warnings = write_report.validation_summary.issues
    for warning in write_report.warnings:
        logger.warning(
            "Reverse pipeline warning: %s route=%s sheet=%s row=%s",
            warning.code,
            warning.route_name,
            warning.sheet_name,
            warning.row_index,
        )

    # 8. Diff validation — assert only writable cells changed
    diff_path = run_dir / "template_diff.json"
    validator = TemplateDiffValidator()
    diff_report = validator.validate(
        template_path=str(template_xlsx),
        submission_path=submission_path,
        template_profile=template_profile,
        write_report=write_report,
        run_id=run_id,
        output_path=str(diff_path),
    )

    # 9. Emit artifacts
    profile_path = run_dir / "template_profile.json"
    plan_path = run_dir / "mapping_plan.json"
    report_path = run_dir / "write_report.json"
    pending_path = run_dir / "pending_review.json"

    profile_path.write_text(template_profile.model_dump_json(indent=2), encoding="utf-8")
    plan_path.write_text(plan.model_dump_json(indent=2), encoding="utf-8")
    report_path.write_text(write_report.model_dump_json(indent=2), encoding="utf-8")
    pending_path.write_text(
        json.dumps([r.model_dump() for r in plan.pending_review], indent=2),
        encoding="utf-8",
    )

    return {
        "run_id": run_id,
        "run_dir": str(run_dir),
        "bid_sheets_profiled": len(template_profile.bid_sheets),
        "export_rows_loaded": len(export_rows),
        "instructions_resolved": len(instructions),
        "cells_written": write_report.cells_written,
        "cells_skipped": write_report.cells_skipped,
        "cells_skipped_preserved": (
            write_report.cells_skipped_preserved
            + write_report.cells_skipped_preserved_resolver
        ),
        "cells_skipped_low_confidence": write_report.cells_skipped_low_confidence,
        "rows_skipped_descriptor_mismatch": write_report.rows_skipped_descriptor_mismatch,
        "duplicate_export_routes": write_report.duplicate_export_routes,
        "auto_write_threshold": auto_write_threshold_resolved,
        "no_bid_lanes": len(no_bid_lanes),
        "no_bid_lane_names": no_bid_lanes,
        "validation_warnings": (
            write_report.validation_summary.issue_counts["warning"]
            if write_report.validation_summary
            else 0
        ),
        "warnings": [warning.model_dump() for warning in write_report.warnings],
        "pending_review_count": len(plan.pending_review),
        "llm_iterations": plan.iterations_run,
        "diff_passed": diff_report.passed,
        "diff_violations": diff_report.unexpected_changes,
        "diff_missing_writes": diff_report.missing_writes,
        "submission": submission_path,
        "write_report": str(report_path),
        "template_profile": str(profile_path),
        "mapping_plan": str(plan_path),
        "pending_review": str(pending_path),
        "template_diff": str(diff_path),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run the Rehydrate (reverse) pipeline locally."
    )
    parser.add_argument(
        "--template",
        required=True,
        help="Path to the original customer template (.xls or .xlsx).",
    )
    parser.add_argument(
        "--export",
        required=True,
        help="Path to the priced export file (FTL LaneExport.xlsx).",
    )
    parser.add_argument(
        "--output-root",
        default="artifacts/local_rehydrate",
        help="Directory where rehydrate artifacts are written.",
    )
    parser.add_argument(
        "--planner-mode",
        default="mock",
        choices=["mock", "live"],
        help="ReversePlanner mode: 'mock' (default) or 'live' (LLM iterative).",
    )
    parser.add_argument(
        "--confidence-threshold",
        type=float,
        default=0.70,
        help="Confidence score below which a mapping is flagged for review (default 0.70).",
    )
    parser.add_argument(
        "--auto-write-threshold",
        type=float,
        default=None,
        help=(
            "Mappings with confidence below this score are NOT auto-written "
            "(falls back to $REHYDRATE_MIN_WRITE_CONFIDENCE, then 0.85)."
        ),
    )
    parser.add_argument(
        "--interactive",
        action="store_true",
        default=False,
        help="Prompt for human override on low-confidence mappings before writing.",
    )
    args = parser.parse_args()

    result = run_rehydrate(
        template_path=args.template,
        export_path=args.export,
        output_root=args.output_root,
        planner_mode=args.planner_mode,
        confidence_threshold=args.confidence_threshold,
        interactive=args.interactive,
        auto_write_threshold=args.auto_write_threshold,
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
