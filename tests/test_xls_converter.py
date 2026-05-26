"""tests/test_xls_converter.py

TD-001 / TD-006 — Tests for xls_converter.py.

Uses only stdlib + openpyxl so no xlrd is needed for the happy path; the
.xls conversion test is skipped when xlrd is not installed (CI without it
will still pass all other tests).
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from src.function_app.services.xls_converter import (
    _fix_zip_value,
    ensure_xlsx,
    is_xls,
)

# ── Helpers ───────────────────────────────────────────────────────────────────

def _xlrd_available() -> bool:
    try:
        import xlrd  # noqa: F401
        return True
    except ImportError:
        return False


def _write_fake_xls(path: Path) -> None:
    """Write a file that starts with OLE2 magic bytes (not a real .xls)."""
    path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 512)


def _write_real_xlsx(path: Path) -> None:
    """Write a minimal .xlsx using openpyxl."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"  # type: ignore[union-attr]
    ws.append(["Name", "ZIP Code", "Value"])  # type: ignore[union-attr]
    ws.append(["Alice", "06851", 42])  # type: ignore[union-attr]
    wb.save(str(path))


# ── is_xls ────────────────────────────────────────────────────────────────────

class TestIsXls(unittest.TestCase):
    def test_detects_ole2_magic_bytes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "fake.xls"
            _write_fake_xls(p)
            self.assertTrue(is_xls(p))

    def test_rejects_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "real.xlsx"
            _write_real_xlsx(p)
            self.assertFalse(is_xls(p))

    def test_rejects_missing_file(self) -> None:
        self.assertFalse(is_xls(Path("/does/not/exist.xls")))

    def test_rejects_short_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "tiny.xls"
            p.write_bytes(b"\xd0\xcf")  # only 2 bytes — not enough
            self.assertFalse(is_xls(p))


# ── _fix_zip_value ────────────────────────────────────────────────────────────

class TestFixZipValue(unittest.TestCase):
    def test_pads_short_zip(self) -> None:
        result = _fix_zip_value(6851.0, "ZIP Code")
        self.assertEqual(result, "06851")

    def test_leaves_normal_zip_as_string(self) -> None:
        result = _fix_zip_value(90248.0, "Postal Code")
        self.assertEqual(result, "90248")

    def test_ignores_non_zip_column(self) -> None:
        result = _fix_zip_value(12345.0, "Annual Volume")
        self.assertEqual(result, 12345.0)

    def test_ignores_non_float(self) -> None:
        result = _fix_zip_value("90210", "ZIP")
        self.assertEqual(result, "90210")

    def test_ignores_fractional_float(self) -> None:
        result = _fix_zip_value(12345.6, "ZIP")
        self.assertEqual(result, 12345.6)

    def test_case_insensitive_column_match(self) -> None:
        result = _fix_zip_value(1234.0, "zip_code")
        self.assertEqual(result, "01234")

    def test_postal_column_name(self) -> None:
        result = _fix_zip_value(2345.0, "postal")
        self.assertEqual(result, "02345")


# ── ensure_xlsx ───────────────────────────────────────────────────────────────

class TestEnsureXlsx(unittest.TestCase):
    def test_passthrough_for_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "file.xlsx"
            _write_real_xlsx(p)
            out, converted = ensure_xlsx(p)
            self.assertFalse(converted)
            self.assertEqual(out, p)

    @unittest.skipUnless(
        _xlrd_available(),
        "xlrd not installed — skipping .xls conversion test",
    )
    def test_converts_xls_to_xlsx(self) -> None:
        fixture = Path(
            "C:/Users/maxbush/OneDrive - Microsoft/Documents/Customers/RXO"
            "/bid_export/source_docs/Original Customer File 1.xls"
        )
        if not fixture.exists():
            self.skipTest("Fixture file not available in this environment")

        out, converted = ensure_xlsx(fixture)
        self.assertTrue(converted)
        self.assertTrue(out.exists())
        self.assertEqual(out.suffix, ".xlsx")

        from openpyxl import load_workbook
        wb = load_workbook(str(out), data_only=True, read_only=True)
        self.assertGreater(len(wb.sheetnames), 0)
        wb.close()


if __name__ == "__main__":
    unittest.main()

