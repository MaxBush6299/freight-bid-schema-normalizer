"""tests/test_rehydrate_pipeline_runner.py

Smoke / integration test for the end-to-end ``run_rehydrate`` pipeline.

Wires ``TemplateProfiler → ReversePlanner (mock) → resolve_instructions
→ TemplateAwareWriter → TemplateDiffValidator`` against the synthetic
Coupa workbook and asserts that:

* every expected artifact file is emitted
* the summary dict carries the documented keys with the right shape
* the mock plan produces non-zero writes (3 matching mappings × 3 routes)
* unmatched template routes show up in ``no_bid_lane_names``
* the diff validator passes on the writer's own output
"""
from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src.function_app.local_rehydrate_runner import (
    apply_plan,
    prepare_plan,
    run_rehydrate,
)
from src.function_app.services.mapping_cache_store import MappingCacheStore
from tests._coupa_fixtures import (
    BID_SHEET_NAME,
    FIRST_DATA_ROW,
    ROUTE_NAMES,
    build_export_workbook,
    build_template_workbook,
)

_EXPECTED_ARTIFACTS = (
    "submission.xlsx",
    "template_profile.json",
    "mapping_plan.json",
    "write_report.json",
    "pending_review.json",
    "template_diff.json",
)


class TestRehydratePipelineRunner(unittest.TestCase):
    """End-to-end smoke test of the rehydrate pipeline."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        # Isolate the mapping cache so tests never share state.
        self._cache_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._cache_dir.cleanup)
        self._old_cache_env = os.environ.get("REHYDRATE_MAPPING_CACHE_ROOT")
        os.environ["REHYDRATE_MAPPING_CACHE_ROOT"] = self._cache_dir.name
        self.addCleanup(self._restore_cache_env)

    def _restore_cache_env(self) -> None:
        if self._old_cache_env is None:
            os.environ.pop("REHYDRATE_MAPPING_CACHE_ROOT", None)
        else:
            os.environ["REHYDRATE_MAPPING_CACHE_ROOT"] = self._old_cache_env

    def test_happy_path_writes_expected_cells_and_artifacts(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        result = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(out_root),
            planner_mode="mock",
        )

        # 1 bid sheet × 3 routes, 3 of 3 mock mappings hit the fixture
        # (Customer Linehaul Rate → Freight \nPrice, Equipment Type
        #  Detail → Equipment Type, Currency → Currency).
        # NOTE: The mock plan intentionally maps Customer Linehaul Rate
        # (not RXO All In Customer Rate) into Freight Price because that
        # template column is linehaul-only — fuel is computed separately
        # by the template's own formula. Writing an all-in value would
        # double-count fuel.
        self.assertEqual(result["bid_sheets_profiled"], 1)
        self.assertEqual(result["export_rows_loaded"], len(ROUTE_NAMES))
        self.assertEqual(result["instructions_resolved"], 9)
        self.assertEqual(result["cells_written"], 9)
        self.assertEqual(result["cells_skipped"], 0)
        self.assertEqual(result["no_bid_lanes"], 0)
        self.assertEqual(result["no_bid_lane_names"], [])
        self.assertEqual(result["validation_warnings"], 0)
        self.assertEqual(result["pending_review_count"], 0)
        self.assertEqual(result["llm_iterations"], 0)
        self.assertTrue(result["diff_passed"])
        self.assertEqual(result["diff_violations"], 0)
        self.assertEqual(result["diff_missing_writes"], 0)

        run_dir = Path(result["run_dir"])
        self.assertTrue(run_dir.exists())
        for fname in _EXPECTED_ARTIFACTS:
            self.assertTrue(
                (run_dir / fname).exists(),
                msg=f"missing artifact: {fname}",
            )

        # Sanity-check the submission .xlsx — Freight \nPrice should be
        # populated in the first slot for every route.
        wb = load_workbook(result["submission"])
        try:
            sheet = wb[BID_SHEET_NAME]
            # Freight \nPrice header lives at col 9 in the fixture
            for row_offset in range(len(ROUTE_NAMES)):
                val = sheet.cell(
                    row=FIRST_DATA_ROW + row_offset,
                    column=9,
                ).value
                self.assertIsNotNone(val, f"row {row_offset}: freight cell is blank")
                self.assertIsInstance(val, (int, float))
        finally:
            wb.close()

    def test_no_bid_lane_when_template_route_missing_from_export(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        # Export covers only the first 2 routes; the third becomes no_bid
        export = build_export_workbook(
            self.tmp / "export.xlsx",
            route_names=ROUTE_NAMES[:-1],
        )

        result = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out"),
            planner_mode="mock",
        )

        # 3 mappings × 2 matched routes
        self.assertEqual(result["cells_written"], 6)
        self.assertEqual(result["no_bid_lanes"], 1)
        self.assertEqual(result["no_bid_lane_names"], [ROUTE_NAMES[-1]])
        self.assertEqual(result["validation_warnings"], 1)
        warning = result["warnings"][0]
        self.assertEqual(warning["code"], "no_bid_lane")
        self.assertEqual(warning["route_name"], ROUTE_NAMES[-1])
        self.assertEqual(warning["sheet_name"], BID_SHEET_NAME)
        # Diff still passes — no_bid lanes just leave cells blank
        self.assertTrue(result["diff_passed"])

    def test_extra_export_route_with_no_template_match_is_ignored(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(
            self.tmp / "export.xlsx",
            extra_route_without_template_match="GHOST-LANE",
        )

        result = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out"),
            planner_mode="mock",
        )

        # 4 export rows loaded but only 3 match → 9 instructions, 0 no_bid
        self.assertEqual(result["export_rows_loaded"], 4)
        self.assertEqual(result["instructions_resolved"], 9)
        self.assertEqual(result["no_bid_lanes"], 0)

    def test_mapping_plan_artifact_records_mock_mode(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")

        result = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out"),
            planner_mode="mock",
        )

        plan_doc = json.loads(Path(result["mapping_plan"]).read_text(encoding="utf-8"))
        self.assertEqual(plan_doc["planner_mode"], "mock")
        self.assertEqual(plan_doc["iterations_run"], 0)
        self.assertGreater(len(plan_doc["mappings"]), 0)

    def test_pipeline_preserves_pre_filled_bid_inputs(self) -> None:
        """If a customer has already filled in slot-0 Currency on the template,
        the pipeline must NOT overwrite that value (this is the regression test
        for the 180-cell overwrite bug observed against the real workbook)."""
        template = build_template_workbook(
            self.tmp / "template.xlsx",
            prefill_descriptors=True,
            prefill_bid_inputs=True,  # slot-0 Currency = "EUR"
        )
        export = build_export_workbook(self.tmp / "export.xlsx")

        result = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out"),
            planner_mode="mock",
            # default auto_write_threshold=0.85; mock confidence=1.0 → eligible
        )

        # 3 mappings × 3 routes = 9 candidate writes.  3 of them (Currency)
        # target a pre-filled cell and must be preserved.
        self.assertEqual(result["instructions_resolved"], 6)
        self.assertEqual(result["cells_written"], 6)
        self.assertEqual(result["cells_skipped_preserved"], 3)
        self.assertTrue(result["diff_passed"])

        # And the on-disk submission must still carry "EUR" in every slot-0
        # Currency cell, not the export's "USD".
        wb = load_workbook(result["submission"])
        try:
            sheet = wb[BID_SHEET_NAME]
            for row_offset in range(len(ROUTE_NAMES)):
                currency_val = sheet.cell(
                    row=FIRST_DATA_ROW + row_offset,
                    column=8,  # slot 0 Currency in the fixture
                ).value
                self.assertEqual(
                    currency_val, "EUR",
                    f"row {row_offset}: pre-filled currency was overwritten",
                )
        finally:
            wb.close()


class TestPrepareAndApplyPlan(unittest.TestCase):
    """The two-step API: prepare_plan + apply_plan, plus cache round-trip."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        # Isolate the cache for this test
        self._cache_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._cache_dir.cleanup)
        self.cache_store = MappingCacheStore(root=self._cache_dir.name)

    def test_prepare_plan_returns_plan_without_writes(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        prepared = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )

        # prepare_plan must NOT touch the output directory.
        self.assertFalse(out_root.exists(), "prepare_plan must not create artifacts")
        # The cache must be empty (no save happened yet).
        self.assertEqual(self.cache_store.list_cached(), [])
        self.assertFalse(prepared.cache_hit)
        self.assertIsNone(prepared.cache_record)

        # Sanity: the plan carries the new mock mappings (Customer Linehaul Rate
        # not RXO All In Customer Rate).
        source_fields = [m.source_field for m in prepared.plan.mappings]
        self.assertIn("Customer Linehaul Rate", source_fields)
        self.assertNotIn("RXO All In Customer Rate", source_fields)

    def test_apply_plan_writes_submission_and_caches_plan(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        prepared = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )

        result = apply_plan(
            prepared,
            output_root=str(out_root),
            approved_by="test-suite",
            cache_store=self.cache_store,
        )

        # Submission written
        self.assertTrue(Path(result["submission"]).is_file())
        self.assertEqual(result["cells_written"], 9)  # 3 mappings × 3 routes
        self.assertTrue(result["diff_passed"])

        # Plan was cached
        self.assertTrue(result["cache_saved"])
        self.assertFalse(result["cache_hit"])
        cache_listing = self.cache_store.list_cached()
        self.assertEqual(len(cache_listing), 1)
        self.assertEqual(
            cache_listing[0], prepared.template_profile.template_fingerprint
        )

        # Cached plan round-trips
        record = self.cache_store.load(cache_listing[0])
        self.assertIsNotNone(record)
        assert record is not None
        self.assertEqual(record.approved_by, "test-suite")
        self.assertEqual(
            record.plan.template_fingerprint,
            prepared.template_profile.template_fingerprint,
        )

    def test_second_prepare_plan_hits_cache(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        # First run — populates the cache.
        first = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        apply_plan(
            first,
            output_root=str(out_root),
            approved_by="alice",
            cache_store=self.cache_store,
        )
        # Mutate the cached plan to prove the second prepare uses it (not the
        # fresh planner output): change the Currency mapping's reasoning.
        record = self.cache_store.load(first.template_profile.template_fingerprint)
        assert record is not None
        plan = record.plan
        for fm in plan.mappings:
            if fm.source_field == "Currency":
                fm.reasoning = "[from cache] operator confirmed"
        self.cache_store.save(plan, approved_by="alice")

        # Second run — should hit the cache.
        second = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        self.assertTrue(second.cache_hit)
        self.assertIsNotNone(second.cache_record)
        assert second.cache_record is not None
        self.assertEqual(second.cache_record.approved_by, "alice")
        # The mutation we made should be visible (proves we got the cached plan).
        currency_mapping = next(
            m for m in second.plan.mappings if m.source_field == "Currency"
        )
        self.assertIn("from cache", (currency_mapping.reasoning or ""))

    def test_prepare_plan_with_use_cache_false_bypasses_cache(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        first = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        apply_plan(
            first,
            output_root=str(out_root),
            approved_by="alice",
            cache_store=self.cache_store,
        )

        second = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            use_cache=False,
            cache_store=self.cache_store,
        )
        self.assertFalse(second.cache_hit)
        self.assertIsNone(second.cache_record)

    def test_apply_plan_with_cache_hit_does_not_re_save(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        out_root = self.tmp / "out"

        first = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        apply_plan(
            first,
            output_root=str(out_root),
            approved_by="alice",
            cache_store=self.cache_store,
        )

        second = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        self.assertTrue(second.cache_hit)
        result = apply_plan(
            second,
            output_root=str(out_root),
            approved_by="bob",  # different operator — but it's a cache hit
            cache_store=self.cache_store,
        )
        # cache hit -> we never overwrite the existing entry on apply
        self.assertFalse(result["cache_saved"])
        record = self.cache_store.load(first.template_profile.template_fingerprint)
        assert record is not None
        self.assertEqual(record.approved_by, "alice")  # not "bob"

    def test_apply_plan_rejects_approved_plan_with_mismatched_fingerprint(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")

        prepared = prepare_plan(
            template_path=str(template),
            export_path=str(export),
            planner_mode="mock",
            cache_store=self.cache_store,
        )
        # Build an alternate plan with a different fingerprint
        forged = prepared.plan.model_copy(deep=True)
        forged.template_fingerprint = "deadbeef00000000"

        with self.assertRaises(ValueError):
            apply_plan(
                prepared,
                output_root=str(self.tmp / "out"),
                approved_plan=forged,
                cache_store=self.cache_store,
            )


class TestCacheIntegrationWithRunRehydrate(unittest.TestCase):
    """run_rehydrate honours the use_cache / save_to_cache args."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self._cache_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._cache_dir.cleanup)
        self._old_env = os.environ.get("REHYDRATE_MAPPING_CACHE_ROOT")
        os.environ["REHYDRATE_MAPPING_CACHE_ROOT"] = self._cache_dir.name
        self.addCleanup(self._restore_env)

    def _restore_env(self) -> None:
        if self._old_env is None:
            os.environ.pop("REHYDRATE_MAPPING_CACHE_ROOT", None)
        else:
            os.environ["REHYDRATE_MAPPING_CACHE_ROOT"] = self._old_env

    def test_run_rehydrate_with_save_to_cache_false_leaves_cache_empty(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")
        run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out"),
            planner_mode="mock",
            save_to_cache=False,
        )
        store = MappingCacheStore()
        self.assertEqual(store.list_cached(), [])

    def test_run_rehydrate_default_uses_and_saves_cache(self) -> None:
        template = build_template_workbook(self.tmp / "template.xlsx")
        export = build_export_workbook(self.tmp / "export.xlsx")

        first = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out1"),
            planner_mode="mock",
        )
        self.assertFalse(first["cache_hit"])
        self.assertTrue(first["cache_saved"])

        second = run_rehydrate(
            template_path=str(template),
            export_path=str(export),
            output_root=str(self.tmp / "out2"),
            planner_mode="mock",
        )
        self.assertTrue(second["cache_hit"])
        self.assertFalse(second["cache_saved"])  # already cached → not re-saved


if __name__ == "__main__":
    unittest.main()
